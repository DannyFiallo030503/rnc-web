"""Aggregate-only image, PDF and spreadsheet renderers."""

import csv
import io
from django.http import HttpResponse


def export_result(fmt, table, payload, params):
    filename = "consulta-rnc-sintetica"
    caveat = (
        "Método candidato; denominadores sintéticos, no son resultados oficiales. "
        "Los conteos de personas por localización pueden solaparse."
    )
    metadata = [
        ("Fuente y corte", payload["note"]),
        ("Período", next((f["value"] for f in payload["filters"] if f["label"] == "Período"), "Todos")),
        ("Territorio", next((f["value"] for f in payload["filters"] if f["label"] == "Territorio"), "Todos")),
        ("Sexo", next((f["value"] for f in payload["filters"] if f["label"] == "Sexo"), "Todos")),
        ("Edad", next((f["value"] for f in payload["filters"] if f["label"] == "Edad"), "Todos")),
        ("Grupo de localización", next((f["value"] for f in payload["filters"] if f["label"] == "Grupo de localización ilustrativo"), "Todos")),
        ("Agrupación", payload["dimension_label"]),
        ("Indicador", payload["metric"]),
        ("Unidad", payload["unit"]),
        ("Gráfico", payload["chart_label"]),
        ("Nota", caveat),
    ]
    header = [
        "Categoría",
        "Tumores",
        "Personas distintas",
        "Población sintética usada como denominador",
        "Valor",
    ]
    values = [
        [r[k] for k in ["label", "tumors", "persons", "population", "value"]]
        for r in table
    ]
    if len(payload["datasets"]) > 1:
        header += [
            d["label"]
            + (" (%)" if params["chart"] == "percent" else " · " + payload["metric"])
            for d in payload["datasets"]
        ]
        values = [
            row + [s["value"] for s in item["series"]]
            for row, item in zip(values, table)
        ]
    if fmt == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Contexto de la consulta", "Selección"])
        writer.writerows(metadata)
        writer.writerow([])
        writer.writerow(header)
        writer.writerows(values)
        response = HttpResponse(
            "\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8"
        )
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        book = Workbook()
        sheet = book.active
        sheet.title = "Consulta"
        sheet.append(["Contexto de la consulta", "Selección"])
        for row in metadata:
            sheet.append(row)
        sheet.append([])
        sheet.append(header)
        for row in values:
            sheet.append(row)
        header_row = len(metadata) + 3
        for c in sheet[header_row]:
            c.font = Font(bold=True)
        from openpyxl.utils import get_column_letter

        for col in range(1, len(header) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 36 if col == 1 else 32
        sheet.freeze_panes = f"A{header_row + 1}"
        output = io.BytesIO()
        book.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    elif fmt in {"png", "svg", "pdf"}:
        # OO Matplotlib avoids global pyplot state in concurrent requests.
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        import textwrap

        figure = Figure(figsize=(12, 9))
        FigureCanvasAgg(figure)
        axes = figure.subplots()
        labels = payload["labels"]
        kind = params["chart"]
        colors = [
            "#177c70",
            "#de765c",
            "#748cb5",
            "#b49a52",
            "#8b77a9",
            "#60a2a2",
            "#bb7792",
            "#7c9561",
        ]
        datasets = payload["datasets"]
        def number_label(value):
            numeric = float(value)
            if numeric.is_integer():
                return f"{int(numeric):,}".replace(",", ".")
            return f"{numeric:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")

        if kind in {"pie", "doughnut"}:
            pie_values = datasets[0]["data"]
            pie_value_iter = iter(pie_values)

            def pie_label(pct):
                value = next(pie_value_iter)
                suffix = "%" if "(%)" in payload["metric"] else ""
                return (
                    f"{number_label(value)}%"
                    if suffix
                    else f"{number_label(value)}\n({pct:.1f}%)"
                )

            _, _, pie_labels = axes.pie(
                pie_values,
                labels=labels,
                autopct=pie_label,
                colors=colors,
                wedgeprops={"width": 0.45} if kind == "doughnut" else None,
            )
            for index, text in enumerate(pie_labels):
                text.set_color("#ffffff" if index == 0 else "#183f39")
        elif kind == "bar":
            bars = axes.barh(labels, datasets[0]["data"], color=colors[0])
            axes.bar_label(bars, labels=[number_label(v) if v else "" for v in datasets[0]["data"]], padding=4)
            axes.margins(x=0.12)
            axes.invert_yaxis()
            axes.set_xlabel(payload["axis_unit"])
            axes.set_ylabel(payload["dimension_label"])
        elif kind == "line":
            axes.plot(labels, datasets[0]["data"], marker="o", color=colors[0])
            axes.margins(y=0.12)
            for x, value in enumerate(datasets[0]["data"]):
                axes.annotate(number_label(value), (x, value), xytext=(0, 8), textcoords="offset points", ha="center")
            axes.set_xlabel(payload["dimension_label"])
            axes.set_ylabel(payload["axis_unit"])
        elif kind in {"grouped", "stacked", "percent"}:
            bottom = [0] * len(labels)
            for index, dataset in enumerate(datasets):
                x = (
                    [i + (-0.2 if index == 0 else 0.2) for i in range(len(labels))]
                    if kind == "grouped"
                    else list(range(len(labels)))
                )
                bars = axes.bar(
                    x,
                    dataset["data"],
                    width=0.4 if kind == "grouped" else 0.75,
                    bottom=None if kind == "grouped" else bottom,
                    color=colors[index],
                    label=dataset["label"],
                )
                shown_values = [
                    (number_label(value) + ("%" if kind == "percent" else ""))
                    if value and (kind != "percent" or value >= 6)
                    else ""
                    for value in dataset["data"]
                ]
                axes.bar_label(
                    bars,
                    labels=shown_values,
                    label_type="center" if kind in {"stacked", "percent"} else "edge",
                    padding=3,
                    fontsize=8,
                    color=("#ffffff" if index == 0 else "#183f39")
                    if kind in {"stacked", "percent"}
                    else "#183f39",
                )
                bottom = [a + b for a, b in zip(bottom, dataset["data"])]
            axes.set_xticks(range(len(labels)), labels, rotation=45, ha="right")
            axes.legend()
            axes.set_xlabel(payload["dimension_label"])
            axes.set_ylabel("Porcentaje de cada categoría (%)" if kind == "percent" else payload["axis_unit"])
            if kind == "percent":
                axes.set_ylim(0, 100)
        else:
            bars = axes.bar(labels, datasets[0]["data"], color=colors[0])
            axes.bar_label(bars, labels=[number_label(v) if v else "" for v in datasets[0]["data"]], padding=4)
            axes.margins(y=0.12)
            axes.tick_params(axis="x", rotation=60)
            axes.set_xlabel(payload["dimension_label"])
            axes.set_ylabel(payload["axis_unit"])
        axes.set_title(payload["dimension_label"])
        figure.suptitle("RNC / Cuba · " + payload["metric"] + " · " + payload["unit"], fontsize=15)
        figure.subplots_adjust(left=0.25, right=0.97, top=0.87, bottom=0.30)
        figure.text(
            0.5, 0.035,
            "\n".join(textwrap.wrap(payload["note"] + " · " + payload["filter_summary"] + " · " + caveat, 150)),
            ha="center", va="bottom", fontsize=7, parse_math=False,
        )
        output = io.BytesIO()
        if fmt == "pdf":
            from matplotlib.backends.backend_pdf import PdfPages

            with PdfPages(output) as pdf:
                pdf.savefig(figure)
                # Detail pages preserve category names, context and tabular values.
                for start in range(0, len(values), 12):
                    detail = Figure(figsize=(11.7, 8.3))
                    FigureCanvasAgg(detail)
                    detail.text(
                        0.06, 0.94, "RNC / Cuba · Detalle de consulta", fontsize=16
                    )
                    detail.text(
                        0.06,
                        0.88,
                        "\n".join(textwrap.wrap(payload["note"] + " · " + payload["filter_summary"] + " · " + caveat, 140)),
                        fontsize=8,
                        va="top",
                        parse_math=False,
                    )
                    ax = detail.add_axes([0.05, 0.1, 0.9, 0.64])
                    ax.axis("off")
                    cells = [
                        [
                            "\n".join(textwrap.wrap(str(v), 28 if i == 0 else 13))
                            for i, v in enumerate(row)
                        ]
                        for row in values[start : start + 12]
                    ]
                    grid = ax.table(
                        cellText=cells,
                        colLabels=["\n".join(textwrap.wrap(h, 16)) for h in header],
                        loc="upper center",
                        cellLoc="left",
                        colWidths=[0.30]
                        + [0.70 / (len(header) - 1)] * (len(header) - 1),
                    )
                    grid.auto_set_font_size(False)
                    grid.set_fontsize(9)
                    detail_row_height = min(0.11, 0.78 / max(len(cells), 1))
                    for (row, col), cell in grid.get_celld().items():
                        cell.set_edgecolor("#d6e3dd")
                        cell.set_height(detail_row_height if row else 0.14)
                        if row == 0:
                            cell.set_facecolor("#176b5d")
                            cell.set_text_props(color="white", weight="bold")
                        else:
                            cell.set_facecolor("#f0f5f1" if row % 2 else "white")
                    detail.text(
                        0.94,
                        0.05,
                        f"Detalle {start//12+1}",
                        ha="right",
                        fontsize=8,
                        color="#60756b",
                    )
                    pdf.savefig(detail)
        else:
            context_text = payload["note"] + " · " + payload["filter_summary"]
            if fmt == "svg":
                import matplotlib

                with matplotlib.rc_context({"svg.fonttype": "none"}):
                    figure.savefig(
                        output,
                        format=fmt,
                        dpi=160,
                        metadata={"Description": context_text},
                    )
            else:
                figure.savefig(output, format=fmt, dpi=160)
        response = HttpResponse(
            output.getvalue(),
            content_type={
                "png": "image/png",
                "svg": "image/svg+xml",
                "pdf": "application/pdf",
            }[fmt],
        )
    else:
        return HttpResponse("Formato no disponible", status=400)
    response["Content-Disposition"] = f'attachment; filename="{filename}.{fmt}"'
    response["Cache-Control"] = "no-store"
    return response
