"""Full UI/backend walkthrough. Disposable DB, private storage and free loopback port.

Never uses the developer's accounts/files or sends real email. Every checkpoint
captures the page and inspects DOM geometry/field errors, then verifies state.
"""

import csv
import email
from email import policy
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import re
import secrets
import shutil
import sys
import tempfile
import threading
from wsgiref.simple_server import WSGIRequestHandler, make_server

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
OUTPUT = ROOT / "docs" / "revision-completa-2026-10-16"


def run():
    # Playwright sync keeps a private event loop in this thread. ORM checkpoints
    # below are sequential, outside browser callbacks, on the disposable DB.
    # This flag belongs only to this review process, never the application server.
    os.environ["DJANGO_ALLOW_ASYNC_UNSAFE"] = "true"
    from django.conf import settings

    with tempfile.TemporaryDirectory(prefix="rnc-full-review-") as temporary:
        workspace = Path(temporary)
        settings.DATABASES = {
            "default": {
                "ENGINE": "django.db.backends.sqlite3",
                "NAME": str(workspace / "review.sqlite3"),
            }
        }
        settings.DEBUG = True
        settings.ALLOWED_HOSTS = ["127.0.0.1", "localhost"]
        settings.PRIVATE_STORAGE_ROOT = workspace / "private"
        settings.EMAIL_BACKEND = "django.core.mail.backends.filebased.EmailBackend"
        settings.EMAIL_FILE_PATH = settings.PRIVATE_STORAGE_ROOT / "mail"
        settings.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]
        import django

        django.setup()
        from django.contrib.auth import get_user_model
        from django.contrib.auth.models import Group
        from django.contrib.staticfiles.handlers import StaticFilesHandler
        from django.core.cache import cache
        from django.core.management import call_command
        from django.core.wsgi import get_wsgi_application
        from django.db import connections
        from public_site.models import (
            Delivery,
            Publication,
            AuditEvent,
            AccountPolicy,
            AccountMail,
        )
        from public_site.services import ingest, transition, delivery_records
        from public_site.testing import uploaded
        from playwright.sync_api import sync_playwright

        call_command("migrate", verbosity=0, interactive=False, stdout=io.StringIO())
        # Expected 400/403 scenarios are recorded, not dumped as traceback noise.
        logging.getLogger("django.request").setLevel(logging.CRITICAL)
        logging.getLogger("django.security.csrf").setLevel(logging.CRITICAL)
        User = get_user_model()
        password = "Review<> '" + secrets.token_urlsafe(16)
        new_password = "New<> '" + secrets.token_urlsafe(16)
        next_password = "Next<> '" + secrets.token_urlsafe(16)
        root = User.objects.create_superuser(
            "qa_root",
            password=password,
            email="root@example.invalid",
            first_name="Elena",
            last_name="García",
        )
        for n in range(15):
            User.objects.create_user(
                f"qa_pagina_{n:02}",
                first_name="Persona",
                last_name="Sintética",
                email=f"pagina{n}@example.invalid",
            )

        class QuietHandler(WSGIRequestHandler):
            def log_message(self, *args):
                pass

        server = make_server(
            "127.0.0.1",
            0,
            StaticFilesHandler(get_wsgi_application()),
            handler_class=QuietHandler,
        )
        base = f"http://127.0.0.1:{server.server_port}"
        settings.PUBLIC_ORIGIN = base
        logging.getLogger("django.request").setLevel(logging.CRITICAL)
        logging.getLogger("django.security.csrf").setLevel(logging.CRITICAL)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        OUTPUT.mkdir(parents=True, exist_ok=True)
        checks = []
        shots = []
        js_errors = []
        resources = []
        server_errors = []
        browser_version = None
        failure = None
        page = None
        current_label = ""
        number = 0
        browser = None

        def check(ok, label):
            if not ok:
                try:
                    page.screenshot(
                        path=str(OUTPUT / f"fallo-{current_label}-{number:03}.png"),
                        full_page=True,
                    )
                except Exception:
                    pass
                raise AssertionError(label)
            checks.append({"viewport": current_label, "check": label})

        def attach(p):
            p.on("pageerror", lambda error: js_errors.append(str(error)))
            p.on("requestfailed",lambda request:resources.append(request.url.split("/static/",1)[-1]) if "/static/" in request.url else None)
            p.on(
                "response",
                lambda response: (
                    resources.append(response.url.split("/static/", 1)[-1])
                    if "/static/" in response.url and response.status >= 400
                    else None
                ),
            )
            p.on(
                "response",
                lambda response: (
                    server_errors.append(response.status)
                    if response.status >= 500
                    else None
                ),
            )

        def capture(name, explanation):
            nonlocal number
            number += 1
            if page.locator("canvas").count():
                canvas=page.locator("canvas").first
                canvas.scroll_into_view_if_needed()
                page.evaluate("()=>{const c=document.querySelector('canvas'),chart=Chart.getChart(c);chart.resize();chart.update('none');}")
                check(canvas.evaluate("c=>{const d=c.getContext('2d').getImageData(0,0,c.width,c.height).data;return d.some((v,i)=>i%4===3&&v>0);}"),name+": el canvas contiene píxeles dibujados")
            check(page.evaluate("()=>Array.from(document.styleSheets).some(s=>s.href&&s.href.includes('components.css')&&s.cssRules.length>50)"),name+": hoja de componentes cargada y aplicada")
            page.evaluate("window.scrollTo(0,0)")
            page.wait_for_timeout(100)
            check(
                page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"),
                f"{name}: página sin desbordamiento horizontal",
            )
            # Inspector evidence at each step: field errors have a local message,
            # an accessible relation, and the red border belongs to the control.
            diagnostics = page.locator(
                "[data-field] input[aria-invalid=true], [data-field] select[aria-invalid=true], [data-field] textarea[aria-invalid=true]"
            ).evaluate_all(
                "elements => elements.map(e=>{ const id=e.id+'_error',m=document.getElementById(id),s=getComputedStyle(e); return {id:e.id,message:!!m&&!m.hidden&&m.textContent.trim().length>0,linked:(e.getAttribute('aria-describedby')||'').split(' ').includes(id),border:s.borderColor}; })"
            )
            for d in diagnostics:
                check(
                    d["message"] and d["linked"],
                    f"{name}: error asociado al campo {d['id']}",
                )
                check(
                    d["border"] == "rgb(165, 53, 42)"
                    or page.locator("#" + d["id"]).get_attribute("type") == "checkbox",
                    f"{name}: borde del control inválido {d['id']}",
                )
            rectangles = page.locator(
                ".brand,.header-nav,.account-nav,.mobile-menu"
            ).evaluate_all(
                "es=>es.filter(e=>e.getClientRects().length&&getComputedStyle(e).display!=='none').map(e=>{const r=e.getBoundingClientRect();return {l:r.left,r:r.right,t:r.top,b:r.bottom};})"
            )
            for i, a in enumerate(rectangles):
                for b in rectangles[i + 1 :]:
                    check(
                        min(a["r"], b["r"]) - max(a["l"], b["l"]) <= 1
                        or min(a["b"], b["b"]) - max(a["t"], b["t"]) <= 1,
                        f"{name}: elementos de cabecera separados",
                    )
            filename = f"{current_label}-{number:03}-{name}.png"
            page.screenshot(path=str(OUTPUT / filename), full_page=True)
            shots.append(
                {
                    "viewport": current_label,
                    "file": filename,
                    "step": name,
                    "explanation": explanation,
                    "fields": diagnostics,
                }
            )
            if number % 8 == 0:
                print(
                    json.dumps(
                        {
                            "viewport": current_label,
                            "captured": len(shots),
                            "checkpoint": name,
                        }
                    ),
                    flush=True,
                )

        def goto(path):
            response = page.goto(base + path)
            page.wait_for_load_state("networkidle")
            return response

        def login(username, secret):
            goto("/acceso/")
            page.locator("#id_username").fill(username)
            page.locator("#id_password").fill(secret)
            page.locator(".login-form button[type=submit]").click()
            page.wait_for_load_state("networkidle")

        def logout():
            if page.viewport_size["width"] <= 760:
                page.locator(".mobile-menu summary").click()
                page.locator(".mobile-menu form button").click()
            else:
                page.locator(".account-menu summary").click()
                page.locator(".account-dropdown form button").click()
            page.wait_for_load_state("networkidle")

        def fill_account(username, address, secret=password, temporary=""):
            for name, value in {
                "username": username,
                "first_name": "María José",
                "last_name": "O’Neill-Pérez",
                "email": address,
                "confirm_password": secret,
                "initial_password": temporary,
            }.items():
                page.locator("#id_" + name).fill(value)
            page.locator("#id_role").select_option("operator")

        def submit():
            page.locator(".portal-form button[type=submit]").click()
            page.wait_for_load_state("networkidle")

        def link_for(user):
            call_command("deliver_account_mail", stdout=io.StringIO())
            check(
                AccountMail.objects.filter(user=user, status="sent").exists(),
                "Worker procesó el correo de la cuenta sintética",
            )
            for path in sorted(
                settings.EMAIL_FILE_PATH.glob("*"),
                key=lambda f: f.stat().st_mtime,
                reverse=True,
            ):
                message = email.message_from_bytes(
                    path.read_bytes(), policy=policy.default
                )
                if user.email in str(message.get("To", "")):
                    body = (
                        message.get_body(preferencelist=("plain",))
                        if message.is_multipart()
                        else message
                    )
                    return re.search(r"http://[^\s]+", body.get_content()).group()
            raise AssertionError("No se encontró el correo local esperado")

        try:
            with sync_playwright() as playwright:
                launch = {"headless": True}
                if os.environ.get("RNC_REVIEW_BROWSER"):
                    launch["executable_path"] = os.environ["RNC_REVIEW_BROWSER"]
                browser = playwright.chromium.launch(**launch)
                browser_version = browser.version
                for index, (width, height, label) in enumerate(
                    [(1440, 1000, "pc"), (1024, 900, "laptop"), (390, 844, "movil")]
                ):
                    current_label = label
                    number = 0
                    cache.clear()
                    delivery_records.cache_clear()
                    baseline = ingest(
                        uploaded(extra=index + 1),
                        f"Corte de revisión {label} 2023–2024",
                        root,
                    )
                    transition(
                        baseline.pk,
                        "publish",
                        root,
                        "Publicación sintética inicial de revisión.",
                    )
                    context = browser.new_context(
                        viewport={"width": width, "height": height},
                        reduced_motion="reduce",
                        locale="es-CU",
                    )
                    page = context.new_page()
                    attach(page)
                    admin_page = page
                    goto("/")
                    check(
                        page.locator(".sidebar-home").count() == 0,
                        "Eliminado acceso redundante Volver al inicio",
                    )
                    capture(
                        "inicio",
                        "Inicio con indicadores del corte publicado; la base de revisión está aislada.",
                    )
                    if width <= 760:
                        page.locator(".mobile-menu summary").click()
                        capture("menu-movil", "Menú de páginas accesible en móvil.")
                        page.keyboard.press("Escape")
                    goto("/acceso/")
                    submit()
                    check(
                        page.locator("#id_username").get_attribute("aria-invalid")
                        == "true",
                        "Login vacío señala el usuario",
                    )
                    capture(
                        "login-vacio",
                        "Campos obligatorios señalados sin borde rojo alrededor de la página.",
                    )
                    page.locator("#id_username").fill("qa_root")
                    page.locator("#id_password").fill("incorrecta")
                    submit()
                    check(
                        not page.url.endswith("/administracion/cortes/"),
                        "Credenciales incorrectas no inician sesión",
                    )
                    capture(
                        "login-incorrecto",
                        "Error de credenciales genérico; no revela si la cuenta existe.",
                    )
                    login("qa_root", password)
                    check(
                        page.url.endswith("/administracion/cortes/"),
                        "Login de soporte abre Administración",
                    )
                    page.locator(".brand").click()
                    check(
                        page.url.endswith("/administracion/cortes/"),
                        "Marca izquierda lleva al panel autorizado",
                    )
                    capture(
                        "administracion",
                        "Panel conectado a entregas, publicación e historial persistentes.",
                    )
                    if width > 760:
                        page.locator(".account-menu summary").click()
                        capture(
                            "menu-cuenta",
                            "Menú de cuenta con rol, herramientas y cierre de sesión.",
                        )
                        page.keyboard.press("Escape")
                    if width > 760:
                        page.locator(".sidebar-toggle").click()
                        page.wait_for_timeout(150)
                        check(
                            abs(
                                page.locator(".site-sidebar").bounding_box()["width"]
                                - 76
                            )
                            < 1,
                            "Sidebar plegado conserva navegación con iconos",
                        )
                        capture(
                            "sidebar-plegado", "Barra plegada sin enlace redundante."
                        )
                        page.locator(".sidebar-toggle").click()
                    goto("/administracion/usuarios/nuevo/")
                    before = User.objects.count()
                    submit()
                    check(
                        User.objects.count() == before,
                        "Formulario vacío no crea registros",
                    )
                    capture(
                        "cuenta-vacia",
                        "Resumen con enlaces y errores junto a los campos obligatorios.",
                    )
                    page.locator("#id_username").press_sequentially("ana<")
                    check(
                        "<" not in page.locator("#id_username").input_value(),
                        "Se impide escribir delimitador no admitido en usuario",
                    )
                    capture(
                        "caracter-bloqueado",
                        "El carácter rechazado no entra en el campo; el mensaje indica el formato admitido.",
                    )
                    page.locator("#id_email").fill("correo@@sin-dominio")
                    page.locator("#id_email").blur()
                    check(
                        page.locator("#id_email_error").is_visible(),
                        "Correo inválido se indica antes de enviar",
                    )
                    capture(
                        "correo-invalido", "Error de formato situado debajo del correo."
                    )
                    page.locator("#id_first_name").fill("María")
                    blocked = page.locator("#id_first_name").evaluate(
                        "el=>{const data=new DataTransfer();data.setData('text/plain','<script>alert(1)</script>');const e=new ClipboardEvent('paste',{clipboardData:data,bubbles:true,cancelable:true});el.dispatchEvent(e);return e.defaultPrevented;}"
                    )
                    check(
                        blocked
                        and page.locator("#id_first_name").input_value() == "María",
                        "El pegado no permitido se rechaza sin truncar el nombre",
                    )
                    capture(
                        "pegado-bloqueado",
                        "Se rechaza el pegado completo y se conserva el valor válido anterior.",
                    )
                    username = f"operador_{label}"
                    address = f"operador.{label}@example.invalid"
                    fill_account(username, address, secret="incorrecta")
                    submit()
                    check(
                        not User.objects.filter(username=username).exists(),
                        "Contraseña de soporte errónea impide guardar",
                    )
                    capture(
                        "confirmacion-incorrecta",
                        "El backend identifica el campo de confirmación; conserva los datos no secretos.",
                    )
                    fill_account("qa_root", address)
                    submit()
                    check(
                        User.objects.count() == before,
                        "Usuario duplicado no inserta otra cuenta",
                    )
                    capture(
                        "usuario-duplicado",
                        "Un nombre de usuario ya registrado se señala en su campo.",
                    )
                    fill_account(username, "root@example.invalid")
                    submit()
                    check(
                        page.locator("#id_email_error").is_visible(),
                        "Correo duplicado devuelto en el campo correcto",
                    )
                    capture(
                        "correo-duplicado",
                        "La unicidad se comprueba en backend y se explica junto al correo.",
                    )
                    fill_account(username, address)
                    page.locator(
                        ".portal-form input[name=csrfmiddlewaretoken]"
                    ).evaluate("el=>el.value='a'.repeat(64)")
                    submit()
                    check(
                        User.objects.count() == before,
                        "Una verificación CSRF fallida no crea la cuenta",
                    )
                    capture(
                        "formulario-caducado",
                        "La protección CSRF muestra una explicación del portal y evita guardar.",
                    )
                    goto("/administracion/usuarios/nuevo/")
                    fill_account(username, address)
                    submit()
                    invited = User.objects.get(username=username)
                    check(
                        invited.groups.filter(name="operator").exists()
                        and not invited.has_usable_password(),
                        "Cuenta invitada: rol persistente, sin contraseña conocida por soporte",
                    )
                    check(
                        AccountMail.objects.filter(
                            user=invited, status="pending"
                        ).exists(),
                        "Invitación encolada en BD",
                    )
                    capture(
                        "usuario-creado",
                        "Alta de cuenta y aviso de invitación, con lista paginada.",
                    )
                    goto("/administracion/usuarios/?q=qa_pagina")
                    page.get_by_role("link", name="Siguiente", exact=True).click()
                    check("page=2" in page.url, "Paginación conserva búsqueda")
                    capture(
                        "usuarios-pagina-2",
                        "Segunda página de una búsqueda de cuentas sintéticas.",
                    )
                    goto(f"/administracion/usuarios/{root.pk}/")
                    page.locator("#id_is_active").uncheck()
                    page.locator("#id_confirm_password").fill(password)
                    submit()
                    root.refresh_from_db()
                    check(
                        root.is_active, "Soporte no puede desactivar su propia cuenta"
                    )
                    capture(
                        "autodesactivacion-rechazada",
                        "Una regla de negocio global tiene resumen claro; no marca campos válidos como incorrectos.",
                    )
                    # Follow a real mail token in a separate unauthenticated browser.
                    link = link_for(invited)
                    user_context = browser.new_context(
                        viewport={"width": width, "height": height},
                        locale="es-CU",
                        reduced_motion="reduce",
                    )
                    page = user_context.new_page()
                    attach(page)
                    page.goto(link)
                    capture(
                        "invitacion",
                        "Enlace del buzón privado abre el formulario para elegir contraseña.",
                    )
                    page.locator("#id_new_password1").fill("123")
                    page.locator("#id_new_password2").fill("123")
                    submit()
                    invited.refresh_from_db()
                    check(
                        not invited.has_usable_password(),
                        "Una clave insuficiente no establece credencial",
                    )
                    capture(
                        "contrasena-insuficiente",
                        "La longitud mínima se comprueba antes de enviar el formulario.",
                    )
                    page.locator("#id_new_password1").fill(new_password)
                    page.locator("#id_new_password2").fill("No-coincide-2026!")
                    submit()
                    invited.refresh_from_db()
                    check(
                        not invited.has_usable_password(),
                        "Contraseñas distintas no activan credencial",
                    )
                    capture(
                        "contrasenas-distintas",
                        "La confirmación muestra el error antes de enviar.",
                    )
                    page.locator("#id_new_password2").fill(new_password)
                    submit()
                    invited.refresh_from_db()
                    check(
                        invited.check_password(new_password)
                        and invited.password != new_password,
                        "La contraseña elegida se guarda como hash",
                    )
                    capture(
                        "invitacion-completa",
                        "Contraseña establecida y confirmación del resultado.",
                    )
                    login(username, new_password)
                    capture(
                        "operador-panel",
                        "Operador accede a cargas con permisos limitados.",
                    )
                    response = goto("/administracion/usuarios/")
                    check(
                        response.status == 403,
                        "Operador no puede acceder a cuentas mediante URL",
                    )
                    capture(
                        "operador-sin-permiso",
                        "Denegación del backend con interfaz del portal.",
                    )
                    logout()
                    page.goto(link)
                    check(
                        page.get_by_text(
                            "El enlace no es válido", exact=False
                        ).is_visible(),
                        "Enlace utilizado deja de ser válido",
                    )
                    capture(
                        "enlace-usado",
                        "Un enlace consumido no permite volver a establecer contraseña.",
                    )
                    goto("/cuenta/recuperar/")
                    submit()
                    capture(
                        "recuperacion-vacia",
                        "El correo requerido se indica en su campo.",
                    )
                    page.locator("#id_email").fill("desconocida@example.invalid")
                    submit()
                    generic = page.locator(".account-panel").inner_text()
                    capture(
                        "recuperacion-desconocida",
                        "El resultado no revela si el correo corresponde a una cuenta.",
                    )
                    goto("/cuenta/recuperar/")
                    page.locator("#id_email").fill(address)
                    submit()
                    check(
                        page.locator(".account-panel").inner_text() == generic,
                        "La recuperación conocida y desconocida tiene la misma respuesta pública",
                    )
                    capture(
                        "recuperacion-solicitada",
                        "Respuesta genérica de recuperación; no expone existencia de cuentas.",
                    )
                    recovery = link_for(invited)
                    page.goto(recovery)
                    page.locator("#id_new_password1").fill(next_password)
                    page.locator("#id_new_password2").fill(next_password)
                    submit()
                    invited.refresh_from_db()
                    check(
                        invited.check_password(next_password),
                        "Recuperación reemplaza el hash de la cuenta correcta",
                    )
                    capture(
                        "recuperacion-completa",
                        "Contraseña recuperada mediante el nuevo mensaje privado.",
                    )
                    login(username, new_password)
                    check(
                        page.locator("[data-error-summary]").is_visible(),
                        "La contraseña anterior ya no inicia sesión",
                    )
                    capture(
                        "clave-anterior-rechazada",
                        "Se rechaza la contraseña anterior después de recuperar acceso.",
                    )
                    login(username, next_password)
                    page = admin_page
                    goto(f"/administracion/usuarios/{invited.pk}/")
                    page.locator("#id_is_active").uncheck()
                    page.locator("#id_role").select_option("publisher")
                    page.locator("#id_confirm_password").fill(password)
                    submit()
                    invited.refresh_from_db()
                    check(
                        not invited.is_active
                        and invited.groups.filter(name="publisher").exists(),
                        "Desactivación y cambio de rol guardados en BD",
                    )
                    capture(
                        "cuenta-desactivada",
                        "Cuenta desactivada sin eliminar su historial.",
                    )
                    page = user_context.pages[0]
                    goto("/administracion/cortes/")
                    check(
                        "/acceso/" in page.url,
                        "Desactivar retira el acceso de una sesión existente",
                    )
                    capture(
                        "sesion-revocada",
                        "Una sesión de la cuenta desactivada ya no abre el panel.",
                    )
                    page = admin_page
                    goto(f"/administracion/usuarios/{invited.pk}/")
                    page.locator("#id_is_active").check()
                    page.locator("#id_role").select_option("operator")
                    page.locator("#id_confirm_password").fill(password)
                    submit()
                    invited.refresh_from_db()
                    check(invited.is_active, "Reactivación persistida")
                    # Temporary credential: known only inside this disposable review.
                    goto("/administracion/usuarios/nuevo/")
                    temporary_username = f"temporal_{label}"
                    fill_account(
                        temporary_username,
                        f"temporal.{label}@example.invalid",
                        temporary=password,
                    )
                    submit()
                    temporary_user = User.objects.get(username=temporary_username)
                    check(
                        AccountPolicy.objects.get(
                            user=temporary_user
                        ).must_change_password,
                        "Cuenta temporal marcada para cambio obligatorio",
                    )
                    page = user_context.pages[0]
                    login(temporary_username, password)
                    check(
                        page.url.endswith("/cuenta/contrasena/"),
                        "La contraseña temporal bloquea navegación hasta cambiarse",
                    )
                    capture(
                        "cambio-obligatorio",
                        "La cuenta temporal debe establecer su contraseña privada antes de utilizar el panel.",
                    )
                    page.locator("#id_old_password").fill("incorrecta")
                    page.locator("#id_new_password1").fill(new_password)
                    page.locator("#id_new_password2").fill(new_password)
                    submit()
                    check(
                        AccountPolicy.objects.get(
                            user=temporary_user
                        ).must_change_password,
                        "Fallo de contraseña actual conserva restricción",
                    )
                    capture(
                        "clave-actual-incorrecta",
                        "Error de contraseña actual devuelto por el backend en su campo.",
                    )
                    page.locator("#id_old_password").fill(password)
                    page.locator("#id_new_password1").fill(new_password)
                    page.locator("#id_new_password2").fill(new_password)
                    submit()
                    temporary_user.refresh_from_db()
                    check(
                        temporary_user.check_password(new_password)
                        and not AccountPolicy.objects.get(
                            user=temporary_user
                        ).must_change_password,
                        "Cambio obligatorio guarda hash y retira restricción",
                    )
                    capture(
                        "cambio-obligatorio-completo",
                        "Cambio confirmado; la cuenta puede continuar al panel.",
                    )
                    # Complete upload/publication flow and inspect filesystem + DB.
                    page = admin_page
                    goto("/administracion/cortes/")
                    previous_deliveries = Delivery.objects.count()
                    submit()
                    check(
                        Delivery.objects.count() == previous_deliveries,
                        "Carga vacía no registra entrega",
                    )
                    capture(
                        "carga-vacia",
                        "Nombre, archivo y confirmación tienen errores individuales.",
                    )
                    page.locator("#id_title").fill("Actualización de revisión")
                    page.locator("#id_synthetic").check()
                    page.locator("#id_file").set_input_files(
                        {
                            "name": "archivo.exe",
                            "mimeType": "application/octet-stream",
                            "buffer": b"not csv",
                        }
                    )
                    submit()
                    check(
                        Delivery.objects.count() == previous_deliveries,
                        "Extensión inválida bloqueada antes de enviar",
                    )
                    capture(
                        "archivo-no-admitido",
                        "El selector muestra que se necesita un CSV.",
                    )
                    page.locator("#id_file").set_input_files(
                        {
                            "name": "grande.csv",
                            "mimeType": "text/csv",
                            "buffer": b"x" * (12 * 1024 * 1024 + 1),
                        }
                    )
                    submit()
                    check(
                        Delivery.objects.count() == previous_deliveries,
                        "Archivo mayor de 12 MB no se envía ni registra",
                    )
                    capture(
                        "archivo-demasiado-grande",
                        "El tamaño máximo se comprueba al seleccionar el archivo y antes de subirlo.",
                    )
                    page.locator("#id_file").set_input_files(
                        {
                            "name": "archivo.csv",
                            "mimeType": "text/csv",
                            "buffer": f"estructura,incorrecta\n{label},2\n".encode(),
                        }
                    )
                    submit()
                    rejected = Delivery.objects.latest("created_at")
                    check(
                        rejected.status == "rejected"
                        and rejected.report["accepted_rows"] == 0,
                        "CSV incompatible queda rechazado sin filas incorporadas",
                    )
                    check(
                        (settings.PRIVATE_STORAGE_ROOT / rejected.storage_key).exists(),
                        "Original rechazado conservado en almacenamiento privado",
                    )
                    capture(
                        "csv-rechazado",
                        "El informe técnico se asocia al archivo; se conserva el nombre del corte.",
                    )
                    with page.expect_download() as downloaded:
                        page.get_by_role(
                            "link", name="Descargar CSV de actualización"
                        ).click()
                    download = downloaded.value
                    path = Path(download.path())
                    raw = path.read_bytes()
                    page.locator("#id_file").set_input_files(
                        {
                            "name": "actualizacion.csv",
                            "mimeType": "text/csv",
                            "buffer": raw,
                        }
                    )
                    page.locator("#id_synthetic").check()
                    submit()
                    loaded = Delivery.objects.latest("created_at")
                    check(
                        loaded.status == "validated"
                        and loaded.row_count == baseline.row_count + 1,
                        "CSV válido se incorpora completo y añade un tumor sintético",
                    )
                    check(
                        loaded.checksum == hashlib.sha256(raw).hexdigest(),
                        "Huella en BD coincide con el archivo subido",
                    )
                    stored = settings.PRIVATE_STORAGE_ROOT / loaded.storage_key
                    check(
                        len(json.loads(stored.read_text())) == loaded.row_count
                        and stored.with_suffix(".csv").read_bytes() == raw,
                        "Representación privada y original coinciden con metadatos",
                    )
                    check(
                        stored.stat().st_mode & 0o777 == 0o600,
                        "Archivo privado con permiso 0600",
                    )
                    capture(
                        "corte-validado",
                        "Detalle real de la entrega: conteos, integridad, responsable y opción de publicación.",
                    )
                    detail = page.url.removeprefix(base)
                    page.locator("#id_reason").fill("corto")
                    submit()
                    check(
                        Publication.objects.get(pk=1).delivery_id == baseline.pk,
                        "Motivo insuficiente no publica",
                    )
                    capture(
                        "motivo-insuficiente",
                        "La justificación exige longitud mínima y el corte activo no cambia.",
                    )
                    page.locator("#id_reason").fill(
                        "Publicación sintética revisada paso a paso."
                    )
                    submit()
                    loaded.refresh_from_db()
                    check(
                        loaded.status == "published"
                        and Publication.objects.get(pk=1).delivery_id == loaded.pk,
                        "Publicación y puntero activo persistidos",
                    )
                    check(
                        AuditEvent.objects.filter(
                            action="Publicación", target=str(loaded.pk), actor=root
                        ).exists(),
                        "Publicación auditada con responsable correcto",
                    )
                    capture(
                        "corte-publicado",
                        "Estado publicado e historial de la decisión.",
                    )
                    goto("/")
                    check(
                        int(
                            page.locator(
                                ".synthetic-summary article strong"
                            ).first.inner_text()
                        )
                        == 38 + index,
                        "Inicio refleja exactamente el nuevo corte del último año",
                    )
                    capture(
                        "inicio-actualizado",
                        "Publicar actualiza las cifras del Inicio mediante el backend.",
                    )
                    goto("/administracion/cortes/")
                    page.locator("#id_title").fill("Reenvío de la misma entrega")
                    page.locator("#id_synthetic").check()
                    page.locator("#id_file").set_input_files(
                        {"name": "reenvio.csv", "mimeType": "text/csv", "buffer": raw}
                    )
                    count = Delivery.objects.count()
                    submit()
                    check(
                        Delivery.objects.count() == count,
                        "Reenviar archivo idéntico no duplica entregas",
                    )
                    capture(
                        "archivo-duplicado",
                        "La huella impide duplicar la entrega y el error aparece junto al archivo.",
                    )
                    # One chosen graph at a time; inspect actual chart data.
                    for kind, dimension in [
                        ("bar", "site"),
                        ("column", "site"),
                        ("line", "period"),
                        ("grouped", "site"),
                        ("stacked", "site"),
                        ("percent", "site"),
                        ("pie", "sex"),
                        ("doughnut", "sex"),
                    ]:
                        goto("/consulta/")
                        page.locator("#id_dimension").select_option(dimension)
                        page.locator("#id_chart").select_option(kind)
                        page.locator("#query-form .apply-filters").click()
                        page.wait_for_load_state("networkidle")
                        payload = json.loads(
                            page.locator("#chart-payload").inner_text()
                        )
                        check(
                            payload["type"] == kind
                            and page.locator("canvas").count() == 1,
                            f"Gráfico {kind}: un solo gráfico y datos correctos",
                        )
                        check(
                            page.evaluate(
                                "Boolean(Chart.getChart(document.getElementById('result-chart')))"
                            ),
                            f"Gráfico {kind}: instancia renderizada",
                        )
                        capture(
                            "grafico-" + kind,
                            "Tipo elegido: "
                            + kind
                            + ". La tabla y el gráfico se generan desde la misma consulta.",
                        )
                    goto("/consulta/")
                    check(
                        page.locator("#id_chart option[value=line]").is_disabled(),
                        "No se ofrece línea para categorías sin orden temporal",
                    )
                    goto("/consulta/?chart=line&dimension=site")
                    check(
                        page.locator("#id_chart_error").is_visible(),
                        "El backend rechaza una combinación forzada por URL",
                    )
                    capture(
                        "grafico-incompatible",
                        "Una combinación manipulada se señala en el selector correspondiente.",
                    )
                    goto("/consulta/?year=2024&dimension=age_group&metric=rate")
                    for fmt, signature in [
                        ("png", b"\x89PNG"),
                        ("svg", b"<?xml"),
                        ("pdf", b"%PDF"),
                        ("csv", b"\xef\xbb\xbf"),
                        ("xlsx", b"PK"),
                    ]:
                        with page.expect_download() as event:
                            page.locator(
                                f'button[form="query-form"][value="{fmt}"]'
                            ).click()
                        download = event.value
                        content = Path(download.path()).read_bytes()
                        check(
                            content.startswith(signature),
                            f"Exportación {fmt}: firma de archivo válida",
                        )
                        if fmt == "csv":
                            text = content.decode("utf-8-sig")
                            check(
                                "2024" in text
                                and "Datos sintéticos" in text
                                and "SYN-P-" not in text,
                                f"Exportación CSV: filtro, procedencia y ausencia de microdatos",
                            )
                        if fmt == "xlsx":
                            from openpyxl import load_workbook

                            book = load_workbook(io.BytesIO(content))
                            check(
                                book.active.max_row == 20,
                                "Excel: 18 grupos etarios y cabeceras guardados",
                            )
                        if fmt == "png":
                            from PIL import Image

                            image = Image.open(io.BytesIO(content))
                            image.verify()
                            check(
                                image.width > 1000,
                                "PNG exportado tiene resolución para compartir",
                            )
                        if fmt == "svg":
                            from xml.etree import ElementTree

                            check(
                                ElementTree.fromstring(content).tag.endswith("svg"),
                                "SVG exportado es XML válido",
                            )
                        if label == "pc":
                            exports = OUTPUT / "exportaciones"
                            exports.mkdir(exist_ok=True)
                            (exports / f"consulta.{fmt}").write_bytes(content)
                        capture(
                            "exportacion-" + fmt,
                            "Descarga "
                            + fmt.upper()
                            + " comprobada; conserva la selección y evita exportar identidades.",
                        )
                    goto(detail)
                    page.locator("#id_reason").fill(
                        "Revocación sintética de la publicación de revisión."
                    )
                    submit()
                    loaded.refresh_from_db()
                    check(
                        loaded.status == "revoked"
                        and Publication.objects.get(pk=1).delivery_id is None,
                        "Revocación retira publicación activa y conserva histórico",
                    )
                    capture(
                        "corte-revocado", "La revocación queda persistida y auditada."
                    )
                    goto("/")
                    check(
                        page.get_by_role(
                            "heading", name="No hay un corte publicado"
                        ).is_visible(),
                        "Inicio deja de mostrar el corte revocado",
                    )
                    capture(
                        "inicio-tras-revocar",
                        "Inicio refleja inmediatamente la ausencia de publicación.",
                    )
                    goto(
                        "/consulta/?year=2024&dimension=age_group&metric=rate&format=csv"
                    )
                    check(
                        page.get_by_role(
                            "heading", name="No hay resultados"
                        ).is_visible(),
                        "No se sirve la exportación del corte revocado desde caché",
                    )
                    capture(
                        "consulta-tras-revocar",
                        "La consulta no filtra ni descarga estadísticas retiradas.",
                    )
                    goto("/administracion/cortes/")
                    logout()
                    goto("/administracion/cortes/")
                    check(
                        "/acceso/" in page.url,
                        "Cerrar sesión retira acceso administrativo",
                    )
                    capture(
                        "sesion-cerrada",
                        "Acceso privado vuelve a exigir autenticación.",
                    )
                    # Additional geometry at narrow and intermediate breakpoints.
                    for extra_width in [320, 760, 768]:
                        page.set_viewport_size({"width": extra_width, "height": 900})
                        goto("/acceso/")
                        check(
                            page.evaluate(
                                "document.documentElement.scrollWidth <= innerWidth+1"
                            ),
                            f"Login sin desbordamiento en {extra_width}px",
                        )
                    user_context.close()
                    context.close()
                    print(
                        json.dumps(
                            {
                                "completed": label,
                                "checks": len(checks),
                                "screenshots": len(shots),
                            }
                        ),
                        flush=True,
                    )
                check(not js_errors, "Sin excepciones JavaScript en todo el recorrido")
                check(not resources, "Sin fallos de carga de recursos estáticos")
                check(not server_errors, "Sin respuestas HTTP 500 en todo el recorrido")
        except Exception as exc:
            failure = str(exc)
            for value in [password, new_password, next_password]:
                failure = failure.replace(value, "[credencial de revisión]")
            failure = re.sub(
                r"/cuenta/restablecer/[^\s\"]+",
                "/cuenta/restablecer/[enlace]/",
                failure,
            )
            failure = re.sub(
                r'value="[A-Za-z0-9]{32,128}"', 'value="[valor sensible]"', failure
            )
            if page:
                try:
                    page.screenshot(path=str(OUTPUT / "fallo.png"), full_page=True)
                except Exception:
                    pass
        finally:
            if browser:
                try:
                    browser.close()
                except Exception:
                    pass
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)
            connections.close_all()
            result = {
                "success": failure is None,
                "browser": browser_version,
                "viewports": [1440, 1024, 390],
                "checks": checks,
                "screenshots": shots,
                "javascript_errors": js_errors,
                "resource_failures": resources,
                "server_errors": server_errors,
                "failure": failure,
                "isolation": "Base SQLite, cuentas, correo y archivos privados temporales; eliminados al terminar. Puerto asignado por el sistema. Sin SMTP real.",
            }
            (OUTPUT / "resultados.json").write_text(
                json.dumps(result, ensure_ascii=False, indent=2)
            )
            write_gallery(result)
        if failure:
            raise SystemExit("Revisión fallida: " + failure)
        print(
            json.dumps(
                {
                    "success": True,
                    "checks": len(checks),
                    "screenshots": len(shots),
                    "output": str(OUTPUT),
                },
                ensure_ascii=False,
            )
        )


def write_gallery(result):
    import html

    cards = []
    for item in result["screenshots"]:
        cards.append(
            f'<article data-width="{item["viewport"]}"><a href="{item["file"]}" target="_blank"><img loading="lazy" src="{item["file"]}" alt="{html.escape(item["step"])}"></a><h2>{html.escape(item["viewport"]+" · "+item["step"])}</h2><p>{html.escape(item["explanation"])}</p></article>'
        )
    page = """<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Evidencia de revisión RNC</title><style>body{font:15px system-ui;background:#f1f4ee;color:#193e32;margin:30px}h1{font-size:32px}header{max-width:1000px;margin-bottom:30px}button{padding:10px 18px;border:1px solid #acbfac;background:white;border-radius:7px;cursor:pointer;margin:4px}.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(300px,1fr));gap:20px}article{background:white;border-radius:12px;padding:14px;border:1px solid #dbe5d9}img{width:100%;height:330px;object-fit:contain;object-position:top;background:#e9eee5}h2{font-size:16px}p{line-height:1.6}a{color:#196652}</style><header><h1>Recorrido completo · RNC Web</h1>"""
    page += (
        f'<p>Resultado: <strong>{"APROBADO" if result["success"] else "INCOMPLETO"}</strong> · {len(result["checks"])} comprobaciones · {len(result["screenshots"])} capturas.</p><p>{html.escape(result["isolation"])}</p><p>Cada paso incluye captura, inspección DOM de geometría y errores de campo, y las comprobaciones de estado correspondientes. Las comprobaciones cubren los escenarios registrados, no todos los errores concebibles ni una certificación de seguridad.</p><a href="resultados.json">Ver resultados detallados</a><div><button data-filter="todos">Todos</button><button data-filter="pc">PC · 1440 px</button><button data-filter="laptop">Laptop · 1024 px</button><button data-filter="movil">Móvil · 390 px</button></div></header><main class="grid">'
        + "".join(cards)
        + '</main><script>document.querySelectorAll("button").forEach(b=>b.onclick=()=>document.querySelectorAll("article").forEach(a=>a.hidden=b.dataset.filter!=="todos"&&a.dataset.width!==b.dataset.filter));</script></html>'
    )
    (OUTPUT / "index.html").write_text(page)
    # Contact sheets make every captured checkpoint available for visual inspection.
    from PIL import Image, ImageOps, ImageDraw

    for label in ["pc", "laptop", "movil"]:
        items = [i for i in result["screenshots"] if i["viewport"] == label]
        for start in range(0, len(items), 9):
            sheet = Image.new("RGB", (1260, 1200), "#e9eee5")
            draw = ImageDraw.Draw(sheet)
            for index, item in enumerate(items[start : start + 9]):
                im = Image.open(OUTPUT / item["file"]).convert("RGB")
                im.thumbnail((400, 350))
                x = 20 + (index % 3) * 420
                y = 30 + (index // 3) * 400
                sheet.paste(im, (x + (400 - im.width) // 2, y))
                draw.text((x, y + 356), item["step"], fill="#173e32")
            sheet.save(OUTPUT / f"contacto-{label}-{start//9+1:02}.jpg")


if __name__ == "__main__":
    run()
